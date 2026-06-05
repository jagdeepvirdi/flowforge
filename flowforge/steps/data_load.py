import csv
import logging
import re
from datetime import date, datetime
from pathlib import Path
from typing import Any

from flowforge.steps.base import BaseStep, StepResult

logger = logging.getLogger(__name__)

_VALID_MODES = {'replace', 'append'}
_VALID_SOURCE_TYPES = {'file', 'query'}

_DATE_RE = re.compile(r'^\d{4}-\d{2}-\d{2}$')
_TS_RE   = re.compile(r'^\d{4}-\d{2}-\d{2}[T ]\d{2}:\d{2}')
_BOOL_WORDS = frozenset({'true', 'false', '1', '0', 'yes', 'no', 't', 'f'})
_TEXT_TYPE = {'postgresql': 'TEXT', 'oracle': 'VARCHAR2(4000)'}


def _is_int(s: str) -> bool:
    try:
        int(s)
        return True
    except ValueError:
        return False


def _is_float(s: str) -> bool:
    try:
        float(s)
        return True
    except ValueError:
        return False


def _infer_python_type(non_null: list, db_type: str) -> str | None:
    """Return SQL type inferred from actual Python types, or None if inconclusive."""
    if all(isinstance(v, bool) for v in non_null):
        return 'BOOLEAN' if db_type == 'postgresql' else 'NUMBER(1)'
    if all(isinstance(v, int) and not isinstance(v, bool) for v in non_null):
        return 'BIGINT' if db_type != 'oracle' else 'NUMBER(18)'
    if all(isinstance(v, float) for v in non_null):
        return 'NUMERIC' if db_type != 'oracle' else 'NUMBER'
    if all(isinstance(v, datetime) for v in non_null):
        return 'TIMESTAMP'
    if all(isinstance(v, date) for v in non_null):
        return 'DATE'
    return None


def _infer_string_type(strs: list[str], db_type: str) -> str:
    """Return SQL type inferred from string-valued data (CSV sources)."""
    if all(s.lower() in _BOOL_WORDS for s in strs):
        return 'BOOLEAN' if db_type == 'postgresql' else 'NUMBER(1)'
    if all(_is_int(s) for s in strs):
        return 'BIGINT' if db_type != 'oracle' else 'NUMBER(18)'
    if all(_is_float(s) for s in strs):
        return 'NUMERIC' if db_type != 'oracle' else 'NUMBER'
    if all(_TS_RE.match(s) for s in strs):
        return 'TIMESTAMP'
    if all(_DATE_RE.match(s) for s in strs):
        return 'DATE'
    return _TEXT_TYPE.get(db_type, 'TEXT')


def _infer_col_type(values: list, db_type: str) -> str:
    """Infer the best SQL column type from a sample of values.

    Priority: bool → int → float → timestamp → date → text.
    Empty/null-only columns default to text.
    """
    non_null = [v for v in values if v is not None and str(v).strip() != '']
    if not non_null:
        return _TEXT_TYPE.get(db_type, 'TEXT')
    python_type = _infer_python_type(non_null, db_type)
    if python_type:
        return python_type
    return _infer_string_type([str(v).strip() for v in non_null], db_type)


class DataLoadStep(BaseStep):
    """
    Bulk-loads data into a target database table.

    Source types:
      file  — CSV or Excel file (supports {{ steps.prev.output_path }})
      query — SQL query executed against any configured source connection

    Target modes:
      replace — TRUNCATE then bulk INSERT
      append  — bulk INSERT only

    Config schema:
    {
      "target_connection_id": "uuid",
      "target_table": "staging.sales_{{ current_month }}",
      "mode": "replace",
      "chunk_size": 1000,
      "column_map": {"SRC_COL": "target_col"},
      "source": {
        "type": "file",
        "file_path": "{{ steps.generate_report.output_path }}",
        "file_format": "csv",       // "csv" | "excel" — inferred from extension if omitted
        "sheet_name": "Sheet1"      // Excel only
      }
    }

    Or with a query source:
    {
      "target_connection_id": "uuid",
      "target_table": "oracle_staging.customers",
      "mode": "replace",
      "source": {
        "type": "query",
        "connection_id": "uuid",
        "query": "SELECT id, name, amount FROM source_table WHERE month = '{{ current_month }}'"
      }
    }
    """

    step_type = 'data_load'

    def run(self, context: dict[str, Any]) -> StepResult:
        from flowforge.engine.context import render

        target_connection_id = self.config.get('target_connection_id', '').strip()
        target_table = render(self.config.get('target_table', ''), context)
        mode = self.config.get('mode', 'append')
        chunk_size = int(self.config.get('chunk_size', 1000))
        column_map: dict[str, str] = self.config.get('column_map', {})
        source_cfg: dict = self.config.get('source', {})
        create_if_missing: bool = bool(self.config.get('create_if_missing', False))

        if not target_connection_id:
            return StepResult(success=False, error='target_connection_id is required')
        if not target_table:
            return StepResult(success=False, error='target_table is required')
        if mode not in _VALID_MODES:
            return StepResult(success=False, error=f"mode must be one of: {', '.join(sorted(_VALID_MODES))}")
        source_type = source_cfg.get('type', '')
        if source_type not in _VALID_SOURCE_TYPES:
            return StepResult(success=False, error=f"source.type must be one of: {', '.join(sorted(_VALID_SOURCE_TYPES))}")

        try:
            columns, rows = self._load_source(source_cfg, context, render)
        except Exception as e:
            logger.exception("DataLoad: failed to read source")
            return StepResult(success=False, error=f"Source read failed: {e}")

        if not rows:
            logger.info("DataLoad: 0 rows from source — nothing to load into %s", target_table)
            return StepResult(success=True, rows_affected=0, logs="Source returned 0 rows — nothing loaded")

        if column_map:
            columns = [column_map.get(c, c) for c in columns]

        try:
            from flowforge.connections.factory import get_connection
            conn = get_connection(target_connection_id)
            with conn:
                created = False
                if create_if_missing and not self._table_exists(conn, target_table):
                    self._create_table(conn, target_table, columns, rows)
                    created = True
                    logger.info("DataLoad: created table %s", target_table)
                total = self._bulk_load(conn, target_table, mode, columns, rows, chunk_size)
        except Exception as e:
            logger.exception("DataLoad: insert into %s failed", target_table)
            return StepResult(success=False, error=str(e))

        created_note = ' (table auto-created)' if create_if_missing and created else ''
        logger.info("DataLoad: %d rows → %s (mode=%s)%s", total, target_table, mode, created_note)
        return StepResult(
            success=True,
            rows_affected=total,
            logs=f"Loaded {total:,} rows into {target_table} (mode={mode}){created_note}",
        )

    # ── source readers ────────────────────────────────────────────────────────

    def _load_source(self, source_cfg: dict, context: dict, render) -> tuple[list[str], list[tuple]]:
        source_type = source_cfg['type']
        if source_type == 'file':
            return self._load_file(source_cfg, context, render)
        return self._load_query(source_cfg, context, render)

    def _load_file(self, source_cfg: dict, context: dict, render) -> tuple[list[str], list[tuple]]:
        file_path = Path(render(source_cfg.get('file_path', ''), context))
        if not file_path.exists():
            raise FileNotFoundError(f"Source file not found: {file_path}")

        fmt = source_cfg.get('file_format', '').lower() or file_path.suffix.lower().lstrip('.')
        if fmt in ('csv', 'txt'):
            return self._read_csv(file_path)
        if fmt in ('xlsx', 'xls', 'excel'):
            return self._read_excel(file_path, source_cfg.get('sheet_name'))
        raise ValueError(f"Unknown file format '{fmt}' for {file_path}. Set source.file_format to 'csv' or 'excel'.")

    def _read_csv(self, path: Path) -> tuple[list[str], list[tuple]]:
        with open(path, newline='', encoding='utf-8-sig') as f:
            reader = csv.reader(f)
            headers = [h.strip() for h in next(reader)]
            rows = [tuple(row) for row in reader if any(row)]
        logger.debug("CSV read: %d rows, %d columns from %s", len(rows), len(headers), path)
        return headers, rows

    def _read_excel(self, path: Path, sheet_name: str | None) -> tuple[list[str], list[tuple]]:
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb[sheet_name] if sheet_name else wb.active
        rows_iter = ws.iter_rows(values_only=True)
        headers = [str(c).strip() for c in next(rows_iter)]
        rows = [tuple(row) for row in rows_iter if any(v is not None for v in row)]
        wb.close()
        logger.debug("Excel read: %d rows, %d columns from %s", len(rows), len(headers), path)
        return headers, rows

    def _load_query(self, source_cfg: dict, context: dict, render) -> tuple[list[str], list[tuple]]:
        from flowforge.connections.factory import get_connection
        connection_id = source_cfg.get('connection_id', '').strip()
        if not connection_id:
            raise ValueError("source.connection_id is required when source.type is 'query'")
        sql = render(source_cfg.get('query', ''), context)
        if not sql.strip():
            raise ValueError("source.query is required when source.type is 'query'")
        conn = get_connection(connection_id)
        with conn:
            rows, columns = conn.execute_query_with_columns(sql)
        logger.debug("Query source: %d rows, %d columns", len(rows), len(columns))
        return columns, rows

    # ── target loader ─────────────────────────────────────────────────────────

    def _table_exists(self, conn, table: str) -> bool:
        """Check table existence via catalog queries.

        Uses information_schema / catalog tables rather than a speculative SELECT
        so that PostgreSQL never enters an aborted-transaction state on a miss.
        """
        db_type = getattr(conn, 'db_type', 'postgresql')
        clean  = table.replace('"', '').replace("'", '')
        parts  = clean.split('.')
        tname  = parts[-1]
        schema = parts[0] if len(parts) > 1 else None

        try:
            if db_type == 'oracle':
                if schema:
                    rows = conn.execute_query(
                        "SELECT 1 FROM all_tables "
                        "WHERE table_name = UPPER(:1) AND owner = UPPER(:2)",
                        (tname, schema),
                    )
                else:
                    rows = conn.execute_query(
                        "SELECT 1 FROM user_tables WHERE table_name = UPPER(:1)",
                        (tname,),
                    )
            else:  # postgresql (and generic fallback)
                rows = conn.execute_query(
                    "SELECT 1 FROM information_schema.tables "
                    "WHERE table_name = %s AND table_schema = %s",
                    (tname, schema or 'public'),
                )
            return bool(rows)
        except Exception:
            return False

    def _create_table(self, conn, table: str, columns: list[str], rows: list[tuple]) -> None:
        """CREATE TABLE with column types inferred from the data being loaded.

        Column names are left unquoted so they match the unquoted references in
        _bulk_load's INSERT statement on both PostgreSQL and Oracle.
        """
        db_type = getattr(conn, 'db_type', 'postgresql')
        sample  = rows[:1000]
        col_defs = []
        for i, col in enumerate(columns):
            col_vals  = [row[i] for row in sample if i < len(row)]
            col_type  = _infer_col_type(col_vals, db_type)
            col_defs.append(f'{col} {col_type}')
        conn.execute_write(f'CREATE TABLE {table} ({", ".join(col_defs)})')
        logger.debug(
            'Created %s with inferred types: %s',
            table,
            ', '.join(f'{c}={d.split()[-1]}' for c, d in zip(columns, col_defs)),
        )

    def _bulk_load(self, conn, table: str, mode: str, columns: list[str], rows: list[tuple], chunk_size: int) -> int:
        if mode == 'replace':
            conn.execute_write(f"TRUNCATE TABLE {table}")  # nosec B608
            logger.debug("Truncated %s", table)

        col_list = ', '.join(columns)
        placeholders = conn.make_placeholders(len(columns))
        insert_sql = f"INSERT INTO {table} ({col_list}) VALUES ({placeholders})"  # nosec B608

        total = 0
        for i in range(0, len(rows), chunk_size):
            chunk = rows[i : i + chunk_size]
            total += conn.execute_many(insert_sql, chunk)
            logger.debug("Inserted rows %d–%d into %s", i + 1, i + len(chunk), table)

        return total
